"""Builds the sample agreement used by tests and the demo seed.

Run: python3 fixtures/make_fixture.py
"""
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Pt, Inches

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("MASTER SERVICES AGREEMENT")
r.bold = True
r.font.size = Pt(14)

p = doc.add_paragraph()
p.add_run("This Master Services Agreement (the “")
p.add_run("Agreement").bold = True
p.add_run("”) is entered into as of ")
h = p.add_run("[EFFECTIVE DATE]")
h.font.highlight_color = WD_COLOR_INDEX.YELLOW
p.add_run(" between Acme Holdings, Inc., a Delaware corporation (“")
p.add_run("Company").bold = True
p.add_run("”), and ")
# A placeholder split across runs with different formatting, as Word often leaves them.
a = p.add_run("[COUNTERPARTY")
a.font.highlight_color = WD_COLOR_INDEX.YELLOW
b = p.add_run(" NAME]")
b.font.highlight_color = WD_COLOR_INDEX.YELLOW
b.bold = True
p.add_run(", a [ENTITY TYPE] (“")
p.add_run("Provider").bold = True
p.add_run("”).")

doc.add_paragraph("Services. Provider will perform the services described in each statement of work.", style="List Number")
doc.add_paragraph("Fees. Company will pay Provider [FEE AMOUNT] within [PAYMENT TERMS] days of receiving a correct invoice.", style="List Number")
doc.add_paragraph("Expenses. Company will reimburse reasonable pre-approved travel expenses at cost.", style="List Number")
doc.add_paragraph("Confidentiality. Each party will protect the other party’s confidential information with reasonable care.", style="List Number")
doc.add_paragraph("Data Protection. Provider will process personal data only on Company’s documented instructions and in line with the Data Processing Addendum.", style="List Number")
doc.add_paragraph("Term. This Agreement starts on the Effective Date and continues for [TERM] unless terminated earlier under this Agreement.", style="List Number")
doc.add_paragraph("Governing Law. This Agreement is governed by the laws of [GOVERNING LAW], without regard to conflict of laws rules.", style="List Number")

doc.add_paragraph("Deliverables", style="Heading 2")
table = doc.add_table(rows=2, cols=3)
table.style = "Table Grid"
for cell, text in zip(table.rows[0].cells, ["Deliverable", "Due date", "Fee"]):
    cell.text = ""
    cell.paragraphs[0].add_run(text).bold = True
for cell, text in zip(table.rows[1].cells, ["[DELIVERABLE]", "[DUE DATE]", "[DELIVERABLE FEE]"]):
    cell.text = text

doc.add_paragraph()
doc.add_paragraph("The parties have signed this Agreement as of the Effective Date.")

sig = doc.add_table(rows=4, cols=2)
labels = [("ACME HOLDINGS, INC.", "[COUNTERPARTY NAME]"),
          ("By: [COMPANY SIGNATURE]", "By: [PROVIDER SIGNATURE]"),
          ("Name: [COMPANY SIGNER]", "Name: [PROVIDER SIGNER]"),
          ("Date: [COMPANY DATE]", "Date: [PROVIDER DATE]")]
for row, (left, right) in zip(sig.rows, labels):
    row.cells[0].text = left
    row.cells[1].text = right
for cell in sig.rows[0].cells:
    cell.paragraphs[0].runs[0].bold = True

for section in doc.sections:
    section.left_margin = section.right_margin = Inches(1)

doc.save("fixtures/services-agreement.docx")
print("wrote fixtures/services-agreement.docx")
